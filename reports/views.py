# reports/views.py
# reports/views.py  — FULL REPLACE
import csv
import io
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal

from dateutil.utils import today
from django.contrib.auth.decorators import login_required
from django.db.models import Q, Sum
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import render
from django.utils.dateparse import parse_date
from django.utils import timezone as dj_tz, cache

from accounts.models import User, Store
from reports.accounting import compute_kpi, compute_debts_total, compute_incoming, daily_kassa_series, \
    monthly_profit_compare, monthly_breakdown
from reports.metrics import _date, kpi_block
from sales.models import Transaction, SellerCommission
from inventory.models import Product

from django.core.cache import cache


# ---- Helperlar (qisqa yo'l) ----

def _parse_date(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except Exception:
        return None

def _period(request):
    today = dj_tz.now().date()
    df = _date(request.GET.get("date_from") or "") or today
    dt = _date(request.GET.get("date_to") or "") or today
    if df > dt: df, dt = dt, df
    return df, dt

def _scope_txns(user, qs):
    if getattr(user, "is_owner", False):
        return qs
    return qs.filter(store_id=user.store_id)

def _sum(qs, field="amount"):
    return qs.aggregate(s=Sum(field))["s"] or Decimal("0")

def _sales_qs(user, df, dt):
    qs = Transaction.objects.filter(
        type="sale", is_void=False, is_approved=True,
        created_at__date__gte=df, created_at__date__lte=dt
    )
    return _scope_txns(user, qs)

def _expense_period_qs(user, df, dt):
    # Period expenses = product IS NULL (tasdiqlangan)
    qs = Transaction.objects.filter(
        type="expense", is_approved=True, product__isnull=True,
        created_at__date__gte=df, created_at__date__lte=dt
    )
    return _scope_txns(user, qs)

def _debt_pay_qs(user, df, dt):
    qs = Transaction.objects.filter(
        type="debt_pay", is_approved=True,
        created_at__date__gte=df, created_at__date__lte=dt
    )
    return _scope_txns(user, qs)

def _cons_payout_qs(user, df, dt):
    qs = Transaction.objects.filter(
        type="consignment_payout", is_approved=True,
        created_at__date__gte=df, created_at__date__lte=dt
    )
    return _scope_txns(user, qs)

def _commissions_qs(user, df, dt):
    qs = (SellerCommission.objects
          .select_related("transaction")
          .filter(is_approved=True,
                  transaction__is_void=False,
                  transaction__created_at__date__gte=df,
                  transaction__created_at__date__lte=dt))
    if not getattr(user, "is_owner", False):
        qs = qs.filter(seller_id=user.id)
    return qs

def _gross_profit(user, df, dt):
    sales = _sales_qs(user, df, dt)
    return _sum(sales, "profit")

def _sales_cash_card(user, df, dt):
    s = _sales_qs(user, df, dt)
    return _sum(s, "cash_amount"), _sum(s, "card_amount")

def _debt_cash_card(user, df, dt):
    d = _debt_pay_qs(user, df, dt)
    return _sum(d, "cash_amount"), _sum(d, "card_amount")

def _kassa(user, df, dt, commissions_total, period_exp, cons_payouts):
    sales_cash, sales_card = _sales_cash_card(user, df, dt)
    debt_cash, debt_card = _debt_cash_card(user, df, dt)
    cash_in = sales_cash + debt_cash
    card_in = sales_card + debt_card
    out = period_exp + cons_payouts + commissions_total  # siz tanlagan siyosat: approve bo‘lishi bilan kassadan ayiramiz
    return cash_in + card_in - out, cash_in, card_in


# ---- VIEWLAR ----

@login_required
def store_report(request):
    df, dt = _period(request)
    # store filter ixtiyoriy (owner uchun)
    store_id = request.GET.get("store_id") if getattr(request.user, "is_owner", False) else request.user.store_id
    # kpi allaqachon seller/owner scope qiladi; store filtri bo'lsa, template dan yuboramiz (linklarda)
    kpi = kpi_block(request.user, df, dt)
    ctx = dict(
        df=df, dt=dt,
        store_id=str(store_id or ""),
        stores=Store.objects.order_by("name") if getattr(request.user, "is_owner", False) else None,
        **kpi
    )
    return render(request, "reports/store_report.html", ctx)


@login_required
def owner_dashboard(request):
    # oxirgi 30 kun default
    today = dj_tz.now().date()
    df = _date(request.GET.get("date_from") or "") or (today - timedelta(days=29))
    dt = _date(request.GET.get("date_to") or "") or today
    kpi = kpi_block(request.user, df, dt)
    ctx = dict(df=df, dt=dt, **kpi)
    return render(request, "reports/owner_dashboard.html", ctx)


def _date_or(s, default):
    try:
        d = parse_date(s or "")
        return d or default
    except Exception:
        return default

@login_required
def profit_overview(request):
    period = (request.GET.get("period") or "day").lower()
    today = dj_tz.now().date()
    if period == "month":
        df = today.replace(day=1)
        dt = today
    elif period == "week":
        df = today - timedelta(days=6)
        dt = today
    else:
        df = today
        dt = today

    store_id = None
    stores = None
    if getattr(request.user, "is_owner", False):
        stores = Store.objects.order_by("name")
        store_id = int(request.GET.get("store_id")) if (request.GET.get("store_id") or "").isdigit() else None

    k = compute_kpi(request.user, df, dt, store_id)
    inc = compute_incoming(request.user, df, dt, store_id)
    debts = compute_debts_total(request.user, df, dt, store_id)

    ctx = dict(
        period=period, start_date=df, end_date=dt,
        stores=stores, store_id=str(store_id or ""),
        # KPI
        total_sales=k.total_sales, gross_profit=k.gross_profit,
        total_expense=k.total_expense, total_commission=k.total_commission,
        total_cons_payouts=k.total_cons_payouts,
        cash_in=k.cash_in, card_in=k.card_in,
        kassa_total=k.kassa_total, net_profit=k.net_profit,
        # Qo‘shimcha ko‘rsatkichlar
        incoming_count=inc.count, incoming_value=inc.all_value,
        debt_total=debts.total, debt_paid=debts.paid,
        debt_cash=debts.cash_paid, debt_card=debts.card_paid,
        debt_balance=debts.balance,
    )
    return render(request, "reports/profit_overview.html", ctx)

@login_required
def daily_cash(request):
    """
    Kunlik '0 dan' kassa ko‘rinishi + net profit, grafik uchun ham mos.
    """
    today = dj_tz.now().date()
    df = _date_or(request.GET.get("date_from"), today - timedelta(days=6))
    dt = _date_or(request.GET.get("date_to"), today)

    store_id = None
    stores = None
    if getattr(request.user, "is_owner", False):
        stores = Store.objects.order_by("name")
        store_id = int(request.GET.get("store_id")) if (request.GET.get("store_id") or "").isdigit() else None

    rows = daily_kassa_series(request.user, df, dt, store_id)
    total_kassa = sum([r["kassa"] for r in rows], start=Decimal("0"))
    total_profit = sum([r["net_profit"] for r in rows], start=Decimal("0"))

    return render(request, "reports/daily_cash.html", {
        "rows": rows, "date_from": df, "date_to": dt,
        "stores": stores, "store_id": str(store_id or ""),
        "sum_kassa": total_kassa, "sum_profit": total_profit,
    })

@login_required
def profit_compare(request):
    """
    Oyma-oy kassa va net profit taqqoslash (admin uchun).
    """
    today = dj_tz.now().date()
    start = today.replace(day=1)
    months = int(request.GET.get("months") or "12")
    store_id = None
    stores = None
    if getattr(request.user, "is_owner", False):
        stores = Store.objects.order_by("name")
        store_id = int(request.GET.get("store_id")) if (request.GET.get("store_id") or "").isdigit() else None

    series = monthly_profit_compare(request.user, start, months, store_id)
    return render(request, "reports/profit_compare.html", {
        "series": series, "months": months,
        "stores": stores, "store_id": str(store_id or ""),
    })

@login_required
def network_report(request):
    period = int((request.GET.get("period") or "30"))
    today = date.today()
    date_from = today - timedelta(days=period)

    qs = Transaction.objects.filter(created_at__date__gte=date_from, created_at__date__lte=today, type="sale")
    if not request.user.is_owner:
        qs = qs.filter(store_id=request.user.store_id)

    cache_key = f"netreport:{'all' if request.user.is_owner else request.user.store_id}:{period}"
    cached = cache.get(cache_key)
    if cached:
        return render(request, "reports/network_report.html", cached)

    # TOP stores (amount bo‘yicha)
    top_stores = list(qs.values("store__name").annotate(s=Sum("amount")).order_by("-s")[:5])

    # Transaction -> product -> brand/model bo‘yicha amount yig‘amiz
    top_brands = list(
        qs.values("product__brand__name").annotate(s=Sum("amount")).order_by("-s")[:5]
    )
    top_models = list(
        qs.values("product__brand__name", "product__model__name").annotate(s=Sum("amount")).order_by("-s")[:5]
    )

    ctx = {
        "period": period,
        "top_stores": top_stores,
        "top_brands": top_brands,
        "top_models": top_models,
        "date_from": date_from,
        "date_to": today,
    }
    cache.set(cache_key, ctx, 300)
    return render(request, "reports/network_report.html", ctx)

@login_required
def sales_log(request):
    if not request.user.is_owner:
        return render(request, "reports/sales_log.html", {"error": _("Only owner.")})

    store_id = request.GET.get("store_id") or ""
    seller_id = request.GET.get("seller_id") or ""
    date_from = request.GET.get("date_from") or ""
    date_to   = request.GET.get("date_to") or ""

    qs = Transaction.objects.select_related("product","store","seller","product__brand","product__model")\
            .filter(type="sale").order_by("-created_at")

    if store_id:
        qs = qs.filter(store_id=store_id)
    if seller_id:
        qs = qs.filter(seller_id=seller_id)

    def _pdate(s):
        try: return datetime.strptime(s, "%Y-%m-%d").date()
        except: return None

    df = _pdate(date_from); dt = _pdate(date_to)
    if df: qs = qs.filter(created_at__date__gte=df)
    if dt: qs = qs.filter(created_at__date__lte=dt)

    total_amount = qs.aggregate(s=Sum("amount"))["s"] or 0
    total_profit = qs.aggregate(s=Sum("profit"))["s"] or 0

    ctx = {
        "rows": qs[:500],  # paginate xohlasangiz qo‘shamiz
        "stores": Store.objects.order_by("name"),
        "sellers": User.objects.filter(role="seller").order_by("username"),
        "store_id": store_id, "seller_id": seller_id,
        "date_from": date_from, "date_to": date_to,
        "total_amount": total_amount, "total_profit": total_profit,
    }
    return render(request, "reports/sales_log.html", ctx)


@login_required
def expenses_log(request):
    if not request.user.is_owner:
        return render(request, "reports/expenses_log.html", {"error": _("Only owner.")})

    store_id = request.GET.get("store_id") or ""
    seller_id = request.GET.get("seller_id") or ""
    date_from = request.GET.get("date_from") or ""
    date_to   = request.GET.get("date_to") or ""
    include_commissions = request.GET.get("include_commissions") == "1"
    tx_exp = Transaction.objects.filter(type="expense", is_approved=True, created_at__date__range=(date_from, today))
    tx_cons = Transaction.objects.filter(type="consignment_payout", is_approved=True,
                                         created_at__date__range=(date_from, today))

    commission_total = SellerCommission.objects.filter(is_approved=True,
                                                       transaction__created_at__date__range=(date_from,
                                                                                             today)).aggregate(
        s=Sum("amount"))["s"] or 0
    commission_paid = SellerCommission.objects.filter(is_approved=True, is_paid=True,
                                                      transaction__created_at__date__range=(date_from,
                                                                                            today)).aggregate(
        s=Sum("amount"))["s"] or 0

    # Faqat expense yozuvlari
    qs = (Transaction.objects
          .select_related("product","store","seller","expense_type","product__brand","product__model")
          .filter(type="expense")
          .order_by("-created_at"))

    if store_id:
        qs = qs.filter(store_id=store_id)
    if seller_id:
        qs = qs.filter(seller_id=seller_id)

    def _pdate(s):
        try: return datetime.strptime(s, "%Y-%m-%d").date()
        except: return None

    df = _pdate(date_from); dt = _pdate(date_to)
    if df: qs = qs.filter(created_at__date__gte=df)
    if dt: qs = qs.filter(created_at__date__lte=dt)

    # Asosiy total (faqat expense)
    total = qs.aggregate(s=Sum("amount"))["s"] or 0

    # Komissiyalarni xohlasak qo‘shamiz (har sotuv uchun $5)
    commission_rows = []
    commission_total = 0
    if include_commissions:
        com_qs = (SellerCommission.objects
                  .select_related("transaction","seller","transaction__store","transaction__product",
                                  "transaction__product__brand","transaction__product__model")
                  .all()
                  .order_by("-transaction__created_at"))

        if store_id:
            com_qs = com_qs.filter(transaction__store_id=store_id)
        if seller_id:
            com_qs = com_qs.filter(seller_id=seller_id)
        if df:
            com_qs = com_qs.filter(transaction__created_at__date__gte=df)
        if dt:
            com_qs = com_qs.filter(transaction__created_at__date__lte=dt)

        # “synthetic” qatorlar: expenses jadvaliga o‘xshash ko‘rinishda
        for c in com_qs[:1000]:  # ko‘p bo‘lsa paginate qilamiz
            row = {
                "created_at": c.transaction.created_at,
                "store": c.transaction.store,
                "seller": c.seller,
                "expense_type_name": _("Commission"),
                "product": c.transaction.product,
                "amount": c.amount,
                "note": _("Seller commission"),
                "is_commission": True,
            }
            commission_rows.append(row)
            commission_total += c.amount or 0

    # Template uchun oddiy ro‘yxat (expense rows)
    expense_rows = []
    for r in qs[:1000]:
        expense_rows.append({
            "created_at": r.created_at,
            "store": r.store,
            "seller": r.seller,
            "expense_type_name": r.expense_type.name if r.expense_type else "-",
            "product": r.product,
            "amount": r.amount,
            "note": r.note or "",
            "is_commission": False,
        })

    # Ikkisini qo‘shib, sanasi bo‘yicha kamayish tartibida beramiz
    all_rows = expense_rows + commission_rows
    all_rows.sort(key=lambda x: x["created_at"], reverse=True)

    ctx = {
        "rows": all_rows,
        "stores": Store.objects.order_by("name"),
        "sellers": User.objects.filter(role="seller").order_by("username"),
        "store_id": store_id, "seller_id": seller_id,
        "date_from": date_from, "date_to": date_to,
        "include_commissions": include_commissions,
        "total": total + (commission_total if include_commissions else 0),
        "expense_total": total,
        "commission_total": commission_total,
    }
    return render(request, "reports/expenses_log.html", ctx)


# ---- Analytics 2.0 (24 oylik) ----
@login_required
def analytics_overview(request):
    """
    24 oylik konsolidatsiyalangan grafiklar (owner va seller scope).
    Filtr: months (default 24), store_id (owner uchun ixtiyoriy).
    """
    today = dj_tz.now().date()
    start = today.replace(day=1)
    months = int(request.GET.get("months") or "24")

    store_id = None
    stores = None
    if getattr(request.user, "is_owner", False):
        stores = Store.objects.order_by("name")
        store_id = int(request.GET.get("store_id")) if (request.GET.get("store_id") or "").isdigit() else None

    series = monthly_breakdown(request.user, start, months, store_id)

    # front uchun soddalashtirilgan massivlar
    labels = [f"{row['year']:04d}-{row['month']:02d}" for row in series]
    net_profit = [float(row["net_profit"] or 0) for row in series]
    kassa_total = [float(row["kassa_total"] or 0) for row in series]
    sales = [float(row["total_sales"] or 0) for row in series]
    gross = [float(row["gross_profit"] or 0) for row in series]
    expense = [float(row["total_expense"] or 0) for row in series]
    comm = [float(row["total_commission"] or 0) for row in series]
    cons = [float(row["total_cons_payouts"] or 0) for row in series]
    cash_in = [float(row["cash_in"] or 0) for row in series]
    card_in = [float(row["card_in"] or 0) for row in series]

    ctx = {
        "labels": labels,
        "series": {
            "net_profit": net_profit,
            "kassa_total": kassa_total,
            "sales": sales,
            "gross": gross,
            "expense": expense,
            "comm": comm,
            "cons": cons,
            "cash_in": cash_in,
            "card_in": card_in,
        },
        "months": months,
        "stores": stores,
        "store_id": str(store_id or ""),
    }
    return render(request, "reports/analytics_overview.html", ctx)

def _range_from_request(request):
    df = parse_date(request.GET.get("date_from") or "")
    dt = parse_date(request.GET.get("date_to") or "")
    return df, dt

def _store_id_from_request(request):
    s = (request.GET.get("store_id") or "").strip()
    return int(s) if s.isdigit() else None

def _tx_base(user):
    # Faqat select_related: .only() YO‘Q (FieldError’ning oldini olamiz)
    qs = (Transaction.objects
          .select_related("store", "seller", "product__brand", "product__model"))
    if not getattr(user, "is_owner", False):
        qs = qs.filter(store_id=user.store_id)
    return qs

def _apply_range(qs, df, dt, date_field="created_at__date"):
    if df: qs = qs.filter(**{f"{date_field}__gte": df})
    if dt: qs = qs.filter(**{f"{date_field}__lte": dt})
    return qs

def _apply_store(qs, store_id):
    return qs.filter(store_id=store_id) if store_id else qs

@login_required
def export_all_bundle(request):
    # faqat owner
    if not getattr(request.user, "is_owner", False):
        return HttpResponseForbidden("Only owner can export.")

    df, dt = _range_from_request(request)
    store_id = _store_id_from_request(request)

    # --- Querylar: .only() YO‘Q ---
    sales = (_apply_store(
                _apply_range(
                    _tx_base(request.user).filter(type="sale", is_void=False), df, dt
                ),
                store_id
             )
             .order_by("created_at"))

    expenses = (_apply_store(
                    _apply_range(
                        _tx_base(request.user)
                        .filter(type="expense")
                        .select_related("expense_type"),  # bu yerda zarur
                        df, dt
                    ),
                    store_id
                )
                .order_by("created_at"))

    commissions = (SellerCommission.objects
                   .select_related("transaction", "seller",
                                   "transaction__store",
                                   "transaction__product__brand",
                                   "transaction__product__model"))
    # owner scope allaqachon, lekin store/date bo‘yicha ham cheklaymiz:
    if df: commissions = commissions.filter(transaction__created_at__date__gte=df)
    if dt: commissions = commissions.filter(transaction__created_at__date__lte=dt)
    if store_id: commissions = commissions.filter(transaction__store_id=store_id)
    commissions = commissions.order_by("transaction__created_at")

    cons_payouts = (_apply_store(
                        _apply_range(
                            _tx_base(request.user).filter(type="consignment_payout"), df, dt
                        ),
                        store_id
                    )
                    .order_by("created_at"))

    debt_pay = (_apply_store(
                    _apply_range(
                        _tx_base(request.user).filter(type="debt_pay"), df, dt
                    ),
                    store_id
                )
                .order_by("created_at"))

    # --- XLSXga urinib ko‘ramiz ---
    buf = io.BytesIO()
    try:
        from openpyxl import Workbook
        wb = Workbook()

        # Sales
        ws = wb.active; ws.title = "Sales"
        ws.append(["id","date","store","seller","product","amount","cash","card","cost","profit"])
        for t in sales:
            prod = t.product
            pname = ""
            if prod:
                b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
            ws.append([t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                       t.store.name if t.store_id else "",
                       t.seller.username if t.seller_id else "",
                       pname,
                       float(t.amount or 0), float(t.cash_amount or 0), float(t.card_amount or 0),
                       float(t.cost or 0), float(t.profit or 0)])

        # Expenses
        ws = wb.create_sheet("Expenses")
        ws.append(["id","date","approved","store","seller","product","expense_type","amount","note"])
        for t in expenses:
            prod = t.product
            pname = ""
            if prod:
                b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
            et = t.expense_type.name if t.expense_type_id else ""
            ws.append([t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                       int(bool(t.is_approved)),
                       t.store.name if t.store_id else "",
                       t.seller.username if t.seller_id else "",
                       pname, et,
                       float(t.amount or 0), (t.note or "")])

        # Commissions
        ws = wb.create_sheet("Commissions")
        ws.append(["id","date","approved","paid","seller","store","product","amount"])
        for c in commissions:
            t = c.transaction
            prod = t.product if t else None
            pname = ""
            if prod:
                b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
            ws.append([c.id,
                       t.created_at.strftime("%Y-%m-%d %H:%M") if t else "",
                       int(bool(c.is_approved)), int(bool(c.is_paid)),
                       c.seller.username if c.seller_id else "",
                       t.store.name if (t and t.store_id) else "",
                       pname, float(c.amount or 0)])

        # Consignment payouts
        ws = wb.create_sheet("ConsignmentPayouts")
        ws.append(["id","date","approved","store","product","amount","note"])
        for t in cons_payouts:
            prod = t.product
            pname = ""
            if prod:
                b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
            ws.append([t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                       int(bool(t.is_approved)),
                       t.store.name if t.store_id else "",
                       pname, float(t.amount or 0), (t.note or "")])

        # Debt payments
        ws = wb.create_sheet("DebtPayments")
        ws.append(["id","date","approved","store","seller","debtor","phone","amount","cash","card"])
        for t in debt_pay:
            ws.append([t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                       int(bool(t.is_approved)),
                       t.store.name if t.store_id else "",
                       t.seller.username if t.seller_id else "",
                       (t.debtor_name or ""), (t.debtor_phone or ""),
                       float(t.amount or 0), float(t.cash_amount or 0), float(t.card_amount or 0)])

        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        # fayl nomiga filtrlarni ham qo‘shish mumkin (istalsa)
        resp["Content-Disposition"] = 'attachment; filename="export_bundle.xlsx"'
        return resp

    except Exception:
        # Fallback: ZIP of CSVs
        zbuf = io.BytesIO()
        with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
            def add_csv(name, header, rows_iter):
                sio = io.StringIO(); w = csv.writer(sio)
                w.writerow(header)
                for row in rows_iter: w.writerow(row)
                z.writestr(name, sio.getvalue())

            # Sales
            def sales_rows():
                for t in sales:
                    prod = t.product
                    pname = ""
                    if prod:
                        b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                        pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
                    yield [
                        t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                        t.store.name if t.store_id else "",
                        t.seller.username if t.seller_id else "",
                        pname, t.amount, t.cash_amount, t.card_amount, t.cost, t.profit
                    ]
            add_csv("sales.csv",
                ["id","date","store","seller","product","amount","cash","card","cost","profit"],
                sales_rows())

            # Expenses
            def exp_rows():
                for t in expenses:
                    prod = t.product
                    pname = ""
                    if prod:
                        b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                        pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
                    et = t.expense_type.name if t.expense_type_id else ""
                    yield [
                        t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                        int(bool(t.is_approved)),
                        t.store.name if t.store_id else "",
                        t.seller.username if t.seller_id else "",
                        pname, et, t.amount, (t.note or "")
                    ]
            add_csv("expenses.csv",
                ["id","date","approved","store","seller","product","expense_type","amount","note"],
                exp_rows())

            # Commissions
            def com_rows():
                for c in commissions:
                    t = c.transaction
                    prod = t.product if t else None
                    pname = ""
                    if prod:
                        b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                        pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
                    yield [
                        c.id,
                        t.created_at.strftime("%Y-%m-%d %H:%M") if t else "",
                        int(bool(c.is_approved)), int(bool(c.is_paid)),
                        c.seller.username if c.seller_id else "",
                        t.store.name if (t and t.store_id) else "",
                        pname, c.amount
                    ]
            add_csv("commissions.csv",
                ["id","date","approved","paid","seller","store","product","amount"],
                com_rows())

            # Consignment payouts
            def cns_rows():
                for t in cons_payouts:
                    prod = t.product
                    pname = ""
                    if prod:
                        b = getattr(prod, "brand", None); m = getattr(prod, "model", None)
                        pname = f"{b.name if b else ''} {m.name if m else ''}".strip()
                    yield [
                        t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                        int(bool(t.is_approved)),
                        t.store.name if t.store_id else "",
                        pname, t.amount, (t.note or "")
                    ]
            add_csv("consignment_payouts.csv",
                ["id","date","approved","store","product","amount","note"],
                cns_rows())

            # Debt payments
            def debt_rows():
                for t in debt_pay:
                    yield [
                        t.id, t.created_at.strftime("%Y-%m-%d %H:%M"),
                        int(bool(t.is_approved)),
                        t.store.name if t.store_id else "",
                        t.seller.username if t.seller_id else "",
                        (t.debtor_name or ""), (t.debtor_phone or ""),
                        t.amount, t.cash_amount, t.card_amount
                    ]
            add_csv("debt_payments.csv",
                ["id","date","approved","store","seller","debtor","phone","amount","cash","card"],
                debt_rows())

        zbuf.seek(0)
        resp = HttpResponse(zbuf.getvalue(), content_type="application/zip")
        resp["Content-Disposition"] = 'attachment; filename="export_bundle.zip"'
        return resp